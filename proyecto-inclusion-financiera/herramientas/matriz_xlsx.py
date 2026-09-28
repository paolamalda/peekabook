# Genera la matriz en Excel (una hoja por categoría + notas) y el mapa de recursos para Moodle.
import json, os, csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
d = json.load(open(os.path.join(BASE, "entregables/matriz/matriz-productos.json")))
cols = [("id","ID",7),("nombre","Producto",32),("costo_mensual","Costo mensual",16),("costos_clave","Costos clave",40),
 ("identificacion","Identificación aceptada",30),("ssn_itin","SSN / ITIN",28),("efectivo","Depósito o pago en efectivo",28),
 ("cobertura","Cobertura",22),("ventaja","Ventaja",30),("cuidado","Cuidado",30),("estado","Estado",14),("fecha","Verificado",12),("fuente","Fuente",45)]
hdr = PatternFill("solid", fgColor="1F5C4A"); amber = PatternFill("solid", fgColor="FFF4D6")
wb = Workbook(); ws = wb.active; ws.title = "Notas"
ws.append(["Matriz comparativa de productos financieros para personas migrantes (California primero)"]); ws["A1"].font = Font(bold=True, size=14)
ws.append([f"Fecha de verificación: {d['fecha_verificacion']}. Filas en amarillo: por confirmar."]); ws.append([])
for n in d["notas"]: ws.append(["• " + n])
ws.column_dimensions["A"].width = 120
def sheet(title, rows):
    s = wb.create_sheet(title[:31]); s.append([c[1] for c in cols])
    for i, c in enumerate(cols, 1):
        cell = s.cell(1, i); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = hdr
        s.column_dimensions[cell.column_letter].width = c[2]
    for r in rows:
        s.append([r[c[0]] for c in cols])
        if r["estado"] != "verificado":
            for cell in s[s.max_row]: cell.fill = amber
    for row in s.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
    s.freeze_panes = "C2"; s.auto_filter.ref = s.dimensions
sheet("Todos", d["productos"])
for cat in dict.fromkeys(r["categoria"] for r in d["productos"]):
    sheet(cat.split(" (")[0], [r for r in d["productos"] if r["categoria"] == cat])
wb.save(os.path.join(BASE, "entregables/matriz/Matriz_comparativa_v1.xlsx"))

# Mapa de recursos para Moodle
wb2 = Workbook(); s = wb2.active; s.title = "Para saber más"
rows = list(csv.reader(open(os.path.join(BASE, "recursos/para_saber_mas_moodle.csv"))))
for r in rows: s.append(r)
for i, w in enumerate([10,12,50,50,28,70,20], 1):
    c = s.cell(1, i); c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr
    s.column_dimensions[c.column_letter].width = w
s.freeze_panes = "B2"; s.auto_filter.ref = s.dimensions
wb2.save(os.path.join(BASE, "recursos/Para_saber_mas_por_leccion.xlsx"))
print("ok")
