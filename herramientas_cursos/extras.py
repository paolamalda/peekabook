# Piezas complementarias que se generan desde las lecciones de cada curso:
#   guiones(): mensajes de WhatsApp y guiones de audio de «Lo esencial», más recordatorios mensuales del compromiso.
#   kit():     kit para facilitadores: una sesión de 60 minutos por módulo.
#   hojas():   herramientas descargables en Excel (presupuesto, deudas, fondo, meta; y negocio si aplica).
import os, re

def partes(bloque):
    """Datos de una lección en formato v3."""
    def uno(p):
        m = re.search(p, bloque, re.M | re.S); return m.group(1).strip() if m else ""
    plain = lambda t: re.sub(r"\{\{([^|}]+)\|[^}]+\}\}", r"\1", t).replace("**", "").strip()
    rec = uno(r"^--- recuerda\n(.*?)\n\n")
    casos = re.findall(r"^### (?:Caso|Case) \d+\. ([^\n]+)\n([^\n]+)\n\? ([^|\n]+)\|\| ([^\n]+)", bloque, re.M)
    return {"gancho": plain(uno(r"^gancho: ([^\n]*)$")), "objetivo": plain(uno(r"^objetivo: ([^\n]*)$")),
            "idea": plain(uno(r"^> \*\*(?:Idea clave|Key idea):\*\* ([^\n]*)$")),
            "recuerda": [plain(x) for x in re.findall(r"^- (.+)$", rec, re.M)],
            "plan": plain(uno(r"^--- plan\n(.*?)(?:\n\n|\n==)")),
            "ponlo": (plain(uno(r"^--- ponlo\n(.*?)\n(?:respuesta|answer):")), plain(uno(r"^--- ponlo\n.*?\n(?:respuesta|answer): ([^\n]*)$"))),
            "casos": [(plain(t), plain(x), plain(q), plain(a)) for t, x, q, a in casos]}


# ---------------------------------------------------------------------------
def guiones(D, CFG, lecciones, EN):
    T = (lambda es, en: en) if EN else (lambda es, en: es)
    L = [f"# {T('Guiones de WhatsApp y audio', 'WhatsApp and audio scripts')} · {CFG['titulo']}", "",
         T("Para el canal de avisos del programa. Cada lección tiene un mensaje corto de WhatsApp y un guion de audio de 40 a 90 segundos basado en «Lo esencial». Al final van los recordatorios mensuales del compromiso. Nunca pidas datos personales por WhatsApp. El enlace de la lección se escribe como «[por definir]» hasta tener la dirección del curso.",
           "For the program's announcement channel. Each lesson has a short WhatsApp message and a 40- to 90-second audio script based on \"The essentials.\" Monthly commitment reminders are at the end. Never ask for personal data over WhatsApp. The lesson link is written as \"[to be defined]\" until you have the course address."), ""]
    link = T("[por definir]", "[to be defined]")
    for mod, titulo in CFG["modulos"].items():
        L += [f"## {titulo}", ""]
        for code, title, b in lecciones(mod):
            p = partes(b)
            cap = lambda x: x[:1].upper() + x[1:]
            p["idea"] = cap(p["idea"]); fin = lambda x: x.rstrip(" .") + "."
            rec = "\n".join(f"• {r}" for r in p["recuerda"][:3])
            wa = f"*{code} · {title}*\n{p['idea']}\n\n{rec}\n\n{T('Tu paso de esta semana', 'Your step this week')}: {p['plan']}\n\n{T('Lección', 'Lesson')}: {link}"
            audio = (T(f"Hola. Hoy hablamos de esto: {fin(title) if not title.endswith('?') else title} ", f"Hi. Today's topic: {fin(title) if not title.endswith('?') else title} ") + p["gancho"] + " " + p["idea"] + " " +
                     T("Recuerda: ", "Remember: ") + " ".join(fin(r) for r in p["recuerda"][:3]) + " " + T("Tu paso de esta semana: ", "Your step this week: ") + fin(p["plan"]) +
                     T(" Nos escuchamos en la próxima lección.", " Talk to you in the next lesson."))
            L += [f"### {code} · {title}", "", "**WhatsApp**", "", "```", wa, "```", "",
                  f"**{T('Audio', 'Audio')}** ({len(audio.split())} {T('palabras, unos', 'words, about')} {round(len(audio.split()) / 2.4)} {T('segundos', 'seconds')})", "", audio, ""]
    meses = T(["Hoy empieza tu meta «[meta]». Aparta [monto] el día que cobres, antes de gastar. Tú puedes.",
               "Llevas un mes con tu meta «[meta]». ¿Ya apartaste este mes? Cada peso cuenta.",
               "Tercer mes de «[meta]». Si un mes no pudiste, no pasa nada: vuelve a empezar en tu próximo pago.",
               "Revisa cuánto llevas en «[meta]» y anótalo en tu plan de una página. Ver el avance motiva.",
               "Cinco meses con «[meta]». ¿Tu apartado es automático? Si no, prográmalo hoy.",
               "Medio año con «[meta]». Cuéntale a tu persona de confianza cómo vas.",
               "Mes 7 de «[meta]». Si recibiste un ingreso extra, piensa en darle una parte a tu meta.",
               "Revisa tus deudas y tu fondo: ¿tu meta «[meta]» sigue siendo la prioridad correcta?",
               "Mes 9 de «[meta]». Tu constancia cuenta: el hábito ya es tuyo.",
               "Faltan dos meses para cerrar el año de «[meta]». ¿Cuánto te falta?",
               "Mes 11 de «[meta]». Celebra tu avance, aunque sea pequeño.",
               "Un año con «[meta]». Actualiza tu plan de una página y ponle nombre a tu siguiente meta."],
              ["Your goal \"[goal]\" starts today. Set aside [amount] on payday, before you spend. You can do this.",
               "One month with your goal \"[goal]\". Did you set money aside this month? Every dollar counts.",
               "Month three of \"[goal]\". If you missed a month, that's OK: start again on your next paycheck.",
               "Check how much you have in \"[goal]\" and write it in your one-page plan. Seeing progress helps.",
               "Five months with \"[goal]\". Is your transfer automatic? If not, set it up today.",
               "Six months with \"[goal]\". Tell your trusted person how it's going.",
               "Month 7 of \"[goal]\". If you got extra income, consider giving part of it to your goal.",
               "Review your debts and your fund: is \"[goal]\" still the right priority?",
               "Month 9 of \"[goal]\". Your consistency counts: the habit is yours now.",
               "Two months left to close the year of \"[goal]\". How much is left?",
               "Month 11 of \"[goal]\". Celebrate your progress, even if it's small.",
               "One year with \"[goal]\". Update your one-page plan and name your next goal."])
    L += [T("## Recordatorios mensuales del compromiso", "## Monthly commitment reminders"), "",
          T("Se envían una vez al mes, el mismo día, con el nombre de la meta que cada persona eligió en su plan de una página. La evidencia muestra que los recordatorios mensuales con la meta ayudan a ahorrar más; un recordatorio extra después de un atraso no ayudó. Si el canal es de difusión, usa la versión general («tu meta»). Incluye siempre: «Responde BAJA para dejar de recibir mensajes».",
            "Send them once a month, on the same day, with the goal name each person chose in their one-page plan. Evidence shows monthly reminders that mention the goal help people save more; an extra reminder after missing a deposit didn't help. On a broadcast channel, use the general version (\"your goal\"). Always include: \"Reply STOP to stop receiving messages.\""), ""]
    L += [f"| {T('Mes', 'Month')} | {T('Mensaje', 'Message')} |", "|---|---|"] + [f"| {i} | {m} |" for i, m in enumerate(meses, 1)]
    L += ["", T("## Avisos de las encuestas", "## Survey notices"), "",
          T("- Al inicio: «Antes de empezar, contesta la encuesta de inicio (5 minutos, anónima): [por definir]».",
            "- At the start: \"Before you begin, answer the start survey (5 minutes, anonymous): [to be defined].\""),
          T("- A los 30 y 90 días de terminar: «¿Cómo vas con tu plan? Contesta el seguimiento (5 minutos, anónimo): [por definir]».",
            "- 30 and 90 days after finishing: \"How's your plan going? Answer the follow-up (5 minutes, anonymous): [to be defined].\"")]
    open(os.path.join(D, "manual", T("guiones_whatsapp_audio.md", "whatsapp_audio_scripts.md")), "w").write("\n".join(L) + "\n")


# ---------------------------------------------------------------------------
AYUDA = {"mx": ["Línea de la Vida: 800 911 2000 (salud mental y adicciones, las 24 horas)", "CONDUSEF: 55 5340 0999 (bancos, créditos, seguros)",
                "088 de la Guardia Nacional (fraudes y delitos en línea)", "911 (emergencias y violencia)"],
         "us": ["988 Suicide & Crisis Lifeline (call or text, English and Spanish)", "CFPB: consumerfinance.gov (complaints about financial companies)",
                "211 (local help with food, rent and bills)", "National Domestic Violence Hotline: 1-800-799-7233", "911 (emergencies)"]}

def kit(D, CFG, lecciones, EN):
    T = (lambda es, en: en) if EN else (lambda es, en: es)
    pais = "us" if CFG.get("encuesta", "mx").startswith("us") else "mx"
    L = [f"# {T('Kit para facilitadores', 'Facilitator kit')} · {CFG['titulo']}", "",
         T("Una sesión de 60 minutos por módulo, presencial o en línea, para grupos de 8 a 20 personas. Todo sale de las lecciones del módulo: ideas clave, casos y prácticas. Se usa después de que el grupo avanzó en el módulo o como repaso.",
           "One 60-minute session per module, in person or online, for groups of 8 to 20 people. Everything comes from the module's lessons: key ideas, cases and practice. Use it after the group has worked through the module or as a review."), "",
         T("## Reglas para toda sesión", "## Rules for every session"), "",
         T("- Nadie comparte datos personales, números de cuenta ni montos reales: usamos ejemplos inventados.", "- No one shares personal data, account numbers or real amounts: we use made-up examples."),
         T("- No recomendamos instituciones ni productos: enseñamos a verificar y comparar.", "- We don't recommend institutions or products: we teach how to verify and compare."),
         T("- No damos asesoría fiscal, legal ni migratoria personalizada: orientamos a dónde acudir.", "- We don't give personalized tax, legal or immigration advice: we point to where to go."),
         T("- Hablar de dinero cuesta: agradece cada participación y no juzgues.", "- Talking about money is hard: thank every contribution and don't judge."), "",
         T("## Señales de estrés y a dónde canalizar", "## Stress signs and where to refer"), "",
         T("Pon atención si alguien falta seguido, se aísla, menciona que dejó de comprar comida o medicinas, que pidió a prestamistas o apps para gastos básicos, que la amenazan por una deuda o que no duerme por el dinero. Habla en privado, sin presionar, y comparte estas opciones:",
           "Watch for someone who misses sessions often, withdraws, mentions skipping food or medicine, borrowing from lenders or apps for basic expenses, threats over a debt, or not sleeping because of money. Talk privately, without pressure, and share these options:"), ""]
    L += [f"- {x}" for x in AYUDA[pais]] + [""]
    for mod, titulo in CFG["modulos"].items():
        les = [(c, t, partes(b)) for c, t, b in lecciones(mod)]
        casos = [(c, x) for c, t, p in les for x in p["casos"][:1]][:3]
        ponlos = [(c, p["ponlo"]) for c, t, p in les if p["ponlo"][0]][:2]
        res = CFG.get("resultados", {}).get(mod, "")
        L += [f"## {titulo}", "", f"**{T('Resultado', 'Outcome')}:** {res}" if res else "", "",
              f"**{T('Materiales', 'Materials')}:** " + T("celular o proyector con el libro del módulo, hojas del libro de apoyo impresas, tarjetas con los casos, plumas.",
                                                          "phone or projector with the module book, printed pages of the support book, case cards, pens."), "",
              f"| {T('Minutos', 'Minutes')} | {T('Actividad', 'Activity')} |", "|---|---|",
              f"| 0–5 | {T('Bienvenida y pregunta de entrada: «¿Qué paso diste esta semana con tu dinero?»', 'Welcome and opening question: “What step did you take with your money this week?”')} |",
              f"| 5–15 | {T('Ideas clave del módulo (abajo). Pide a alguien que explique una con sus palabras.', 'Key ideas of the module (below). Ask someone to explain one in their own words.')} |",
              f"| 15–35 | {T('Casos en grupos de 3 o 4 (abajo). Cada grupo decide y explica por qué.', 'Cases in groups of 3 or 4 (below). Each group decides and explains why.')} |",
              f"| 35–50 | {T('Práctica: resuelvan los ejercicios (abajo) y comparen.', 'Practice: solve the exercises (below) and compare.')} |",
              f"| 50–57 | {T('Mi compromiso: cada persona escribe un paso con fecha y se lo dice a su pareja de grupo.', 'My commitment: each person writes one step with a date and tells their group partner.')} |",
              f"| 57–60 | {T('Cierre: recuerda la autoevaluación del módulo y el siguiente encuentro.', 'Close: remind them of the module self-assessment and the next meeting.')} |", "",
              f"**{T('Ideas clave', 'Key ideas')}**", ""]
        L += [f"- {c} · {t}: {p['idea']}" for c, t, p in les] + ["", f"**{T('Casos', 'Cases')}**", ""]
        for c, (ct, cx, cq, ca) in casos:
            L += [f"- **{c} · {ct}.** {cx} *{cq}* — {T('Respuesta esperada', 'Expected answer')}: {ca}"]
        L += ["", f"**{T('Práctica', 'Practice')}**", ""]
        for c, (q, a) in ponlos:
            L += [f"- **{c}.** {q} — {T('Respuesta', 'Answer')}: {a}"]
        L += [""]
    open(os.path.join(D, "manual", T("kit_facilitadores.md", "facilitator_kit.md")), "w").write("\n".join(L) + "\n")


# ---------------------------------------------------------------------------
def hojas(D, CFG, EN):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    T = (lambda es, en: en) if EN else (lambda es, en: es)
    usd = CFG.get("encuesta", "mx").startswith("us")
    fmt = '"$"#,##0.00' if usd else '"$"#,##0.00'
    AZ, RO, GR = "0A3161", "E4007C", "E6ECF5"
    H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor=AZ); IN = PatternFill("solid", fgColor="FFF4FA")
    B = Border(bottom=Side(style="thin", color="C9D3E3"))
    wb = Workbook(); first = True

    def hoja(nombre, titulo, nota):
        nonlocal first
        ws = wb.active if first else wb.create_sheet()
        first = False; ws.title = nombre[:31]
        ws["A1"] = titulo; ws["A1"].font = Font(bold=True, size=14, color=AZ)
        ws["A2"] = nota; ws["A2"].font = Font(italic=True, color="5A6478")
        ws.column_dimensions["A"].width = 38
        for col in "BCDEFGH": ws.column_dimensions[col].width = 16
        return ws

    def cab(ws, fila, cols):
        for i, c in enumerate(cols):
            cell = ws.cell(fila, i + 1, c); cell.font = H; cell.fill = HF; cell.alignment = Alignment(wrap_text=True)

    def entrada(cell, val=None):
        if val is not None: cell.value = val
        cell.fill = IN; cell.number_format = fmt

    nota = T("Escribe solo en las celdas rosas. Usa montos inventados o aproximados; este archivo es tuyo y no se comparte.",
             "Type only in the pink cells. Use made-up or approximate amounts; this file is yours and isn't shared.")
    # Presupuesto
    ws = hoja(T("Presupuesto", "Budget"), T("Mi presupuesto del mes", "My monthly budget"), nota)
    cab(ws, 4, [T("Concepto", "Item"), T("Planeado", "Planned"), T("Real", "Actual"), T("Diferencia", "Difference")])
    filas = [T("Lo que me llega (ingreso neto)", "What I take home (net income)"), T("Ahorro, apartado primero", "Savings, set aside first"),
             T("Renta o vivienda", "Rent or housing"), T("Comida", "Food"), T("Transporte", "Transportation"), T("Luz, agua, gas, internet", "Utilities and internet"),
             T("Pagos de deudas", "Debt payments"), T("Escuela e hijos", "School and children"), T("Salud", "Health"), T("Apoyo a la familia", "Family support"),
             T("Gastos del año (dividido entre 12)", "Yearly expenses (divided by 12)"), T("Gustos", "Treats"), T("Otros", "Other")]
    for i, f in enumerate(filas, 5):
        ws.cell(i, 1, f).border = B
        entrada(ws.cell(i, 2)); entrada(ws.cell(i, 3))
        ws.cell(i, 4, f"=C{i}-B{i}").number_format = fmt
    n = 5 + len(filas)
    ws.cell(n + 1, 1, T("Me queda (ingreso menos ahorro y gastos)", "Left over (income minus savings and expenses)")).font = Font(bold=True)
    for col in "BC":
        c = ws[f"{col}{n + 1}"]; c.value = f"={col}5-SUM({col}6:{col}{n - 1})"; c.number_format = fmt; c.font = Font(bold=True, color=RO)
    # Deudas
    ws = hoja(T("Deudas", "Debts"), T("Mi lista de deudas", "My debt list"), nota)
    cab(ws, 4, [T("Deuda", "Debt"), T("Saldo", "Balance"), T("Tasa anual %", "Annual rate %"), T("Pago mínimo", "Minimum payment"),
                T("Orden avalancha", "Avalanche order"), T("Orden bola de nieve", "Snowball order")])
    for i in range(5, 13):
        ws.cell(i, 1).fill = IN
        entrada(ws.cell(i, 2)); ws.cell(i, 3).fill = IN; entrada(ws.cell(i, 4))
        ws.cell(i, 5, f'=IF(C{i}="","",RANK(C{i},$C$5:$C$12,0))')
        ws.cell(i, 6, f'=IF(B{i}="","",RANK(B{i},$B$5:$B$12,1))')
    ws.cell(14, 1, T("Total que debo", "Total owed")).font = Font(bold=True); ws["B14"] = "=SUM(B5:B12)"; ws["B14"].number_format = fmt
    ws.cell(15, 1, T("Total de pagos mínimos al mes", "Total minimum payments per month")).font = Font(bold=True); ws["D15"] = "=SUM(D5:D12)"; ws["D15"].number_format = fmt
    ws.cell(17, 1, T("Avalancha: paga primero la tasa más alta (orden 1). Bola de nieve: paga primero el saldo más chico (orden 1). En las demás, paga el mínimo.",
                     "Avalanche: pay the highest rate first (order 1). Snowball: pay the smallest balance first (order 1). Pay the minimum on the rest."))
    # Fondo
    ws = hoja(T("Fondo", "Fund"), T("Mi fondo de emergencia", "My emergency fund"), nota)
    rows = [(T("Gastos básicos de un mes", "Basic expenses for one month"), None, True), (T("Meses que quiero cubrir", "Months I want to cover"), 3, True),
            (T("Mi meta de fondo", "My fund goal"), "=B4*B5", False), (T("Lo que ya tengo", "What I already have"), None, True),
            (T("Lo que aparto cada mes", "What I set aside each month"), None, True),
            (T("Meses para lograrlo", "Months to reach it"), '=IF(B8>0,ROUNDUP((B6-B7)/B8,0),"")', False)]
    for i, (lab, val, inp) in enumerate(rows, 4):
        ws.cell(i, 1, lab)
        c = ws.cell(i, 2, val)
        if inp: entrada(c, val)
        else: c.number_format = fmt if i == 6 else "0"; c.font = Font(bold=True, color=RO)
    ws["B5"].number_format = "0"
    # Meta con interés compuesto
    ws = hoja(T("Meta", "Goal"), T("Mi meta con interés compuesto", "My goal with compound interest"), nota)
    rows = [(T("Nombre de mi meta", "My goal's name"), T("Mi fondo", "My fund")), (T("Lo que aparto cada mes", "Monthly deposit"), 500),
            (T("Tasa anual estimada %", "Estimated annual rate %"), 7), (T("Años", "Years"), 5),
            (T("Tendré al final (aprox.)", "I'll have at the end (approx.)"), "=IF(B6=0,B5*B7*12,FV(B6/100/12,B7*12,-B5,0))"),
            (T("De eso, lo que aparté", "Of that, what I deposited"), "=B5*B7*12"), (T("Y lo que ganó el interés", "And what interest earned"), "=B8-B9")]
    for i, (lab, val) in enumerate(rows, 4):
        ws.cell(i, 1, lab); c = ws.cell(i, 2, val)
        if i <= 7: c.fill = IN
        if i in (5, 8, 9, 10): c.number_format = fmt
        if i >= 8: c.font = Font(bold=True, color=RO)
    ws["A12"] = T("La tasa es un supuesto para planear; la tasa real puede subir o bajar. Compárala con la inflación.",
                  "The rate is an assumption for planning; the real rate can go up or down. Compare it with inflation.")
    if (CFG.get("negocio") or CFG.get("hojas_negocio")):
        ws = hoja(T("Precio", "Price"), T("Costo y precio de mi producto", "Cost and price of my product"), nota)
        cab(ws, 4, [T("Insumo o costo por unidad", "Supply or cost per unit"), T("Costo", "Cost")])
        for i in range(5, 12): ws.cell(i, 1).fill = IN; entrada(ws.cell(i, 2))
        ws["A13"] = T("Costo total por unidad", "Total cost per unit"); ws["B13"] = "=SUM(B5:B11)"; ws["B13"].number_format = fmt
        ws["A14"] = T("Margen que quiero %", "Margin I want %"); ws["B14"] = 35; ws["B14"].fill = IN
        ws["A15"] = T("Precio sugerido", "Suggested price"); ws["B15"] = "=IF(B14>=100,\"\",B13/(1-B14/100))"; ws["B15"].number_format = fmt; ws["B15"].font = Font(bold=True, color=RO)
        ws["A16"] = T("Ganancia por unidad", "Profit per unit"); ws["B16"] = "=B15-B13"; ws["B16"].number_format = fmt
        ws = hoja(T("Equilibrio", "Break-even"), T("Mi punto de equilibrio", "My break-even point"), nota)
        rows = [(T("Costos fijos del mes (incluye tu sueldo)", "Fixed costs per month (include your salary)"), None), (T("Precio por unidad", "Price per unit"), None),
                (T("Costo variable por unidad", "Variable cost per unit"), None),
                (T("Unidades que necesito vender al mes", "Units I need to sell per month"), '=IF(B6-B7>0,ROUNDUP(B5/(B6-B7),0),"")'),
                (T("Al día (26 días)", "Per day (26 days)"), '=IF(B8="","",ROUNDUP(B8/26,0))')]
        for i, (lab, val) in enumerate(rows, 5):
            ws.cell(i, 1, lab); c = ws.cell(i, 2, val)
            if val is None: entrada(c)
            else: c.font = Font(bold=True, color=RO)
        ws = hoja(T("Flujo", "Cash flow"), T("Mi flujo de 8 semanas", "My 8-week cash flow"), nota)
        cab(ws, 4, [T("Concepto", "Item")] + [T(f"Sem {i}", f"Wk {i}") for i in range(1, 9)])
        ws.cell(5, 1, T("Saldo inicial", "Starting balance")); entrada(ws.cell(5, 2))
        ws.cell(6, 1, T("Entradas (cobros)", "Money in (collections)")); ws.cell(7, 1, T("Salidas (pagos)", "Money out (payments)"))
        ws.cell(8, 1, T("Saldo final", "Ending balance")).font = Font(bold=True)
        for j in range(2, 10):
            col = ws.cell(4, j).column_letter
            entrada(ws.cell(6, j)); entrada(ws.cell(7, j))
            if j > 2: ws.cell(5, j, f"={ws.cell(8, j - 1).column_letter}8").number_format = fmt
            c = ws.cell(8, j, f"={col}5+{col}6-{col}7"); c.number_format = fmt; c.font = Font(bold=True, color=RO)
    # Hojas propias de cada público (curso.json: "hojas": [...])
    def tabla(ws, fila, cols, n, formulas=None, anchos=None):
        cab(ws, fila, cols)
        for r in range(fila + 1, fila + 1 + n):
            for j in range(1, len(cols) + 1):
                f = (formulas or {}).get(j)
                if f: c = ws.cell(r, j, f.format(r=r)); c.number_format = fmt; c.font = Font(bold=True, color=RO)
                elif j == 1: ws.cell(r, j).fill = IN
                else: entrada(ws.cell(r, j))
        return fila + 1 + n
    def total(ws, fila, etiqueta, formula, col=2):
        ws.cell(fila, 1, etiqueta).font = Font(bold=True)
        c = ws.cell(fila, col, formula); c.number_format = fmt; c.font = Font(bold=True, color=RO)
    for h in CFG.get("hojas", []):
        if h == "quincena_turnos":
            ws = hoja(T("Quincena", "Pay period"), T("Mi quincena con turnos extra", "My pay period with extra shifts"), nota)
            rows = [(T("Sueldo base de la quincena", "Base pay for the period"), None), (T("Turnos extra", "Extra shifts"), None), (T("Pago por turno extra", "Pay per extra shift"), None),
                    (T("Descuentos (adelantos, préstamos, FONACOT)", "Deductions (advances, loans)"), None),
                    (T("Lo que me llega", "What I take home"), "=B4+B5*B6-B7"), (T("Ahorro, apartado primero (10%)", "Savings, set aside first (10%)"), "=ROUND(B8*0.1,0)"),
                    (T("Para gastar la quincena", "To spend this period"), "=B8-B9")]
            for i, (lab, val) in enumerate(rows, 4):
                ws.cell(i, 1, lab); c = ws.cell(i, 2, val)
                if val is None: entrada(c)
                else: c.number_format = fmt; c.font = Font(bold=True, color=RO)
            ws["B5"].number_format = "0"
        elif h == "pago_por_dia":
            ws = hoja(T("Mis casas", "My jobs"), T("Lo que gano por día en cada casa", "What I earn per day at each job"), nota)
            fin = tabla(ws, 4, [T("Casa (sin nombres, usa un apodo)", "Job (use a nickname)"), T("Días por semana", "Days per week"), T("Pago por día", "Pay per day"), T("Pasajes que me dan", "Fare they give me"), T("A la semana", "Per week"), T("Al mes (×4.33)", "Per month (×4.33)")],
                        6, {5: "=B{r}*(C{r}+D{r})", 6: "=ROUND(E{r}*4.33,0)"})
            total(ws, fin + 1, T("Total a la semana", "Total per week"), f"=SUM(E5:E{fin - 1})", 5); total(ws, fin + 2, T("Total al mes", "Total per month"), f"=SUM(F5:F{fin - 1})", 6)
            ws.cell(fin + 4, 1, T("Presupuesta con el mes más flojo, no con el mejor.", "Budget with your slowest month, not your best."))
        elif h == "ingreso_variable":
            ws = hoja(T("Ingreso variable", "Variable income"), T("Mis ingresos de 12 meses y mi fondo de sequía", "My 12 months of income and my dry-spell fund"), nota)
            cab(ws, 4, [T("Mes", "Month"), T("Ingreso", "Income")])
            for i in range(12): ws.cell(5 + i, 1, i + 1); entrada(ws.cell(5 + i, 2))
            total(ws, 18, T("Promedio al mes", "Monthly average"), "=IFERROR(AVERAGE(B5:B16),0)"); total(ws, 19, T("Mes más bajo", "Lowest month"), "=MIN(B5:B16)")
            ws.cell(20, 1, T("Gasto básico al mes", "Basic monthly expenses")); entrada(ws.cell(20, 2))
            ws.cell(21, 1, T("Meses de sequía que quiero cubrir", "Dry-spell months to cover")); ws.cell(21, 2, 4).fill = IN
            total(ws, 22, T("Mi meta de fondo de sequía", "My dry-spell fund goal"), "=B20*B21")
        elif h == "bienes":
            ws = hoja(T("Mis bienes", "My assets"), T("Lo que tengo, a nombre de quién y quién lo recibe", "What I own, whose name it's in and who receives it"), T("No anotes números de cuenta ni contraseñas; solo dónde está cada documento.", "Don't write account numbers or passwords; only where each document is."))
            cab(ws, 4, [T("Bien o cuenta", "Asset or account"), T("A nombre de", "In whose name"), T("Valor aproximado", "Approximate value"), T("Documento", "Document"), T("Dónde está el documento", "Where the document is"), T("Beneficiario o heredero", "Beneficiary or heir")])
            for r in range(5, 17):
                for j in range(1, 7):
                    if j == 3: entrada(ws.cell(r, j))
                    else: ws.cell(r, j).fill = IN
            total(ws, 18, T("Valor total aproximado", "Approximate total value"), "=SUM(C5:C16)", 3)
            for col in "BDEF": ws.column_dimensions[col].width = 24
        elif h == "envios":
            ws = hoja(T("Envíos", "Remittances"), T("Lo que cuesta cada envío", "What each transfer costs"), nota)
            fin = tabla(ws, 4, [T("Servicio (apodo)", "Service (nickname)"), T("Envío en dólares", "Amount sent (USD)"), T("Comisión", "Fee"), T("Tipo de cambio que te dan", "Exchange rate you get"), T("Tipo de cambio del día", "Market rate"), T("Pesos que llegan", "Pesos received"), T("Costo total en dólares", "Total cost (USD)")],
                        6, {6: "=B{r}*D{r}", 7: "=IFERROR(C{r}+B{r}*(E{r}-D{r})/E{r},\"\")"})
            ws.cell(fin + 1, 1, T("Compara el costo total: la comisión más el tipo de cambio.", "Compare total cost: the fee plus the exchange rate."))
        elif h == "adelantos":
            ws = hoja(T("Adelantos", "Advances"), T("Adelantos a mi equipo", "Advances to my team"), nota)
            fin = tabla(ws, 4, [T("Persona (apodo)", "Person (nickname)"), T("Adelanto", "Advance"), T("Descuento por pago", "Deduction per pay"), T("Pagos hechos", "Payments made"), T("Falta", "Remaining"), T("Pagos que faltan", "Payments left")],
                        8, {5: "=MAX(B{r}-C{r}*D{r},0)", 6: "=IF(C{r}>0,ROUNDUP(E{r}/C{r},0),\"\")"})
            ws.cell(fin + 1, 1, T("Fondo de adelantos del mes", "Monthly advances fund")); entrada(ws.cell(fin + 1, 2))
            total(ws, fin + 2, T("Me queda en el fondo", "Left in the fund"), f"=B{fin + 1}-SUM(B5:B{fin - 1})")
            for r in range(5, fin): ws.cell(r, 4).number_format = "0"; ws.cell(r, 6).number_format = "0"
        elif h == "semana_apps":
            ws = hoja(T("Mi semana", "My week"), T("Lo que me queda de verdad en la semana", "What I really keep each week"), nota)
            cab(ws, 4, [T("Concepto", "Item"), T("Semana", "Week")])
            ent = [T("Ganancias en la app 1", "Earnings app 1"), T("Ganancias en la app 2", "Earnings app 2"), T("Propinas", "Tips")]
            sal = [T("Gasolina o carga", "Gas or charging"), T("Mantenimiento (aparta por semana)", "Maintenance (set aside weekly)"), T("Renta del vehículo o pago del crédito", "Vehicle rent or loan payment"),
                   T("Plan de datos del celular", "Phone data plan"), T("Seguro del vehículo (parte de la semana)", "Vehicle insurance (weekly share)"), T("Comida en la calle", "Food on the road")]
            r = 5
            for e in ent: ws.cell(r, 1, e); entrada(ws.cell(r, 2)); r += 1
            total(ws, r, T("Entró en la semana", "Money in this week"), f"=SUM(B5:B{r - 1})"); ri = r; r += 1
            s0 = r
            for e in sal: ws.cell(r, 1, e); entrada(ws.cell(r, 2)); r += 1
            total(ws, r, T("Costos de trabajar", "Cost of working"), f"=SUM(B{s0}:B{r - 1})"); rc = r; r += 1
            total(ws, r, T("Me queda de verdad", "What I really keep"), f"=B{ri}-B{rc}"); rq = r; r += 1
            ws.cell(r, 1, T("Horas conectado en la semana", "Hours logged in this week")); ws.cell(r, 2).fill = IN; rh = r; r += 1
            total(ws, r, T("Me queda por hora", "What I keep per hour"), f"=IF(B{rh}>0,B{rq}/B{rh},\"\")"); r += 1
            total(ws, r, T("Apartado para imprevistos y retiro (10%)", "Set aside for emergencies and retirement (10%)"), f"=ROUND(B{rq}*0.1,0)")
        elif h == "pension":
            ws = hoja(T("Mi pensión", "My pension"), T("Mis pensiones y mi mes", "My pensions and my month"), nota)
            rows = [(T("Pensión Bienestar (bimestral)", "Pensión Bienestar (every two months)"), None), (T("Pensión IMSS o ISSSTE (mensual)", "IMSS or ISSSTE pension (monthly)"), None),
                    (T("Apoyo de la familia (mensual)", "Family support (monthly)"), None), (T("Otro ingreso (mensual)", "Other income (monthly)"), None),
                    (T("Lo que me entra al mes", "My monthly income"), "=B4/2+B5+B6+B7"), (T("Descuento por préstamo a cuenta de pensión", "Pension loan deduction"), None),
                    (T("Gastos del mes", "Monthly expenses"), None), (T("Me queda al mes", "Left each month"), "=B8-B9-B10"),
                    (T("Tope de descuento (30% de la pensión IMSS)", "Deduction cap (30% of IMSS pension)"), "=ROUND(B5*0.3,0)")]
            for i, (lab, val) in enumerate(rows, 4):
                ws.cell(i, 1, lab); c = ws.cell(i, 2, val)
                if val is None: entrada(c)
                else: c.number_format = fmt; c.font = Font(bold=True, color=RO)
            ws["A14"] = T("La Pensión Bienestar llega cada dos meses: aparta la mitad para el mes sin depósito.", "Pensión Bienestar arrives every two months: set half aside for the month without a deposit.")
        elif h == "temporada":
            ws = hoja(T("Temporada", "Season"), T("Mi temporada: lo que gano, lo que envío y lo que traigo", "My season: what I earn, send and bring home"), nota)
            rows = [(T("Semanas de contrato", "Contract weeks"), None), (T("Horas por semana (promedio)", "Hours per week (average)"), None), (T("Pago por hora (dólares)", "Pay per hour (USD)"), None),
                    (T("Lo que gano en la temporada", "What I earn in the season"), "=B4*B5*B6"), (T("Gastos allá por semana (comida, teléfono)", "Weekly costs there (food, phone)"), None),
                    (T("Envíos a mi familia por semana", "Weekly remittances to my family"), None), (T("Lo que me queda para traer", "What's left to bring home"), "=B7-B4*(B8+B9)"),
                    (T("Mi meta al regresar", "My goal when I return"), None), (T("Me falta o me sobra", "Short or over"), "=B10-B11")]
            for i, (lab, val) in enumerate(rows, 4):
                ws.cell(i, 1, lab); c = ws.cell(i, 2, val)
                if val is None: entrada(c)
                else: c.number_format = fmt; c.font = Font(bold=True, color=RO)
            for a in ("B4", "B5"): ws[a].number_format = "0"
            ws["B6"].number_format = '"$"#,##0.00'
            ws["A14"] = T("Con la garantía de tres cuartas partes, te deben pagar al menos 75% de las horas del contrato.", "Under the three-fourths guarantee, you must be paid for at least 75% of the contract hours.")
        elif h == "regreso":
            ws = hoja(T("Mis 90 días", "My 90 days"), T("Mis primeros 90 días de regreso", "My first 90 days back"), nota)
            fin = tabla(ws, 4, [T("Concepto", "Item"), T("Mes 1", "Month 1"), T("Mes 2", "Month 2"), T("Mes 3", "Month 3"), T("Total", "Total")], 10, {5: "=SUM(B{r}:D{r})"})
            etiquetas = [T("Lo que traigo o me llega", "What I bring or receive"), T("Ingreso de trabajo", "Work income"), T("Renta o aporte a la casa", "Rent or household share"), T("Comida", "Food"),
                         T("Transporte", "Transportation"), T("Trámites (acta, INE, CURP, cuenta)", "Paperwork (birth certificate, ID, account)"), T("Salud", "Health"), T("Herramientas o capacitación", "Tools or training"),
                         T("Apoyo a la familia", "Family support"), T("Ahorro", "Savings")]
            for i, e in enumerate(etiquetas): ws.cell(5 + i, 1, e)
            total(ws, fin + 1, T("Me queda al final de los 90 días", "Left at the end of 90 days"), f"=SUM(E5:E6)-SUM(E7:E{fin - 1})", 5)
    sig = CFG.get("sigla") or os.path.basename(D)
    d = os.path.join(D, "moodle", "herramientas"); os.makedirs(d, exist_ok=True)
    nombre = T(f"Herramientas_{sig}.xlsx", f"Tools_{sig}.xlsx")
    for f in os.listdir(d): os.remove(os.path.join(d, f))
    wb.save(os.path.join(d, nombre))
    return nombre, wb.sheetnames
